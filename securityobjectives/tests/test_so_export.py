import zipfile
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest
from django.contrib.admin.models import LogEntry
from django.core import mail
from django.urls import reverse
from kombu.exceptions import OperationalError

from governanceplatform.globals import EXPORT
from governanceplatform.models import RegulatorUser, Sector
from securityobjectives.helpers import can_export_security_objectives
from securityobjectives.models import SecurityObjectiveExport
from securityobjectives.tasks import cleanup_so_export_file, generate_so_export_task


def grant_export_right(user):
    RegulatorUser.objects.filter(user=user, regulator=user.regulators.first()).update(can_export_security_objectives=True)


def get_user(populate_so_db, email):
    return next(u for u in populate_so_db["users"] if u.email == email)


def load_export_workbook(export):
    """The stored file is an opaque uuid with no extension, so openpyxl gets a handle."""
    with open(export.get_file_path(), "rb") as export_file:
        return openpyxl.load_workbook(export_file)


def selectable_sectors(declaration):
    """The picker offers leaf sectors only, as the dashboard filter does."""
    return declaration.sectors.filter(parent__isnull=False)


def build_filters(populate_so_db, file_format="xlsx", sectors=None):
    standard = populate_so_db["so_standard"][0]
    declaration = populate_so_db["sas"][0]
    return {
        "regulation": str(standard.regulation_id),
        "standards": [str(standard.id)],
        "years": [str(declaration.year_of_submission)],
        "sectors": sectors if sectors is not None else [str(sector.id) for sector in selectable_sectors(declaration)],
        "statuses": [declaration.status],
        "file_format": file_format,
    }


def run_export(user, filters, settings, tmp_path):
    """Run the task in-process, as CI has no broker to hand it to."""
    settings.PATH_FOR_REPORTING_PDF = str(tmp_path)
    export = SecurityObjectiveExport.objects.create(
        user=user,
        regulation_id=filters["regulation"],
        filename=f"SO_export.{'xlsx' if filters['file_format'] == 'xlsx' else 'zip'}",
    )
    with patch("securityobjectives.tasks.cleanup_so_export_file.apply_async"):
        generate_so_export_task(export.id, user.id, filters, "en")
    export.refresh_from_db()
    return export


@pytest.mark.django_db
def test_right_is_denied_by_default(populate_so_db):
    regulator_admin = get_user(populate_so_db, "regadmin@reg1.lu")

    assert can_export_security_objectives(regulator_admin) is False


@pytest.mark.django_db
def test_right_is_granted_to_a_regulator_user(populate_so_db):
    regulator_user = get_user(populate_so_db, "reguser@reg1.lu")
    grant_export_right(regulator_user)

    assert can_export_security_objectives(regulator_user) is True


@pytest.mark.django_db
def test_operators_and_observers_are_out_of_scope(populate_so_db):
    for email in ["opadmin@com1.lu", "obsadmin@obs1.lu"]:
        user = next((u for u in populate_so_db["users"] if u.email == email), None)
        if user:
            assert can_export_security_objectives(user) is False


@pytest.mark.django_db
def test_dashboard_hides_the_export_link_without_the_right(otp_client, populate_so_db):
    client = otp_client(get_user(populate_so_db, "regadmin@reg1.lu"))

    response = client.get(reverse("securityobjectives"))

    assert response.status_code == 200
    assert "export_security_objectives" not in response.content.decode()


@pytest.mark.django_db
def test_dashboard_shows_the_export_link_with_the_right(otp_client, populate_so_db):
    regulator_admin = get_user(populate_so_db, "regadmin@reg1.lu")
    grant_export_right(regulator_admin)
    client = otp_client(regulator_admin)

    response = client.get(reverse("securityobjectives"))

    assert response.status_code == 200
    assert "export_security_objectives" in response.content.decode()


@pytest.mark.django_db
def test_export_view_is_forbidden_without_the_right(otp_client, populate_so_db):
    client = otp_client(get_user(populate_so_db, "regadmin@reg1.lu"))

    response = client.get(reverse("export_security_objectives"))

    assert response.status_code == 403


@pytest.mark.django_db
def test_export_modal_renders_every_filter(otp_client, populate_so_db):
    regulator_admin = get_user(populate_so_db, "regadmin@reg1.lu")
    grant_export_right(regulator_admin)
    client = otp_client(regulator_admin)

    response = client.get(reverse("export_security_objectives"))
    content = response.content.decode()

    assert response.status_code == 200
    assert 'id="id_regulation"' in content
    assert 'id="id_file_format"' in content
    # The checkbox widget names its select after the field, and export_security_objectives.js
    # selects on exactly these ids and on the plugin's class.
    for field in ["standards", "years", "sectors", "statuses"]:
        assert f'id="{field}"' in content
    assert content.count("multiselectcheckbox") == 4
    assert "so_standards_by_regulation" in content


@pytest.mark.django_db
def test_export_view_dispatches_the_task(otp_client, populate_so_db):
    regulator_admin = get_user(populate_so_db, "regadmin@reg1.lu")
    grant_export_right(regulator_admin)
    client = otp_client(regulator_admin)
    filters = build_filters(populate_so_db)

    with patch("securityobjectives.views.generate_so_export_task.delay") as delay:
        delay.return_value.id = "task-id"
        response = client.post(reverse("export_security_objectives"), filters)

    assert response.status_code == 200
    export = SecurityObjectiveExport.objects.get(id=response.json()["export_id"])
    assert export.user == regulator_admin
    assert export.task_status == "RUNNING"
    delay.assert_called_once()


@pytest.mark.django_db
def test_export_view_reports_a_dead_broker(otp_client, populate_so_db):
    regulator_admin = get_user(populate_so_db, "regadmin@reg1.lu")
    grant_export_right(regulator_admin)
    client = otp_client(regulator_admin)

    with patch(
        "securityobjectives.views.generate_so_export_task.delay",
        side_effect=OperationalError("broker is down"),
    ):
        response = client.post(reverse("export_security_objectives"), build_filters(populate_so_db))

    assert response.status_code == 400
    # The job never ran, so it must not be left behind at RUNNING.
    assert not SecurityObjectiveExport.objects.exists()


@pytest.mark.django_db
def test_status_fails_a_running_job_when_the_stack_is_down(otp_client, populate_so_db):
    regulator_admin = get_user(populate_so_db, "regadmin@reg1.lu")
    grant_export_right(regulator_admin)
    export = SecurityObjectiveExport.objects.create(
        user=regulator_admin,
        regulation_id=populate_so_db["so_standard"][0].regulation_id,
        filename="SO_export.xlsx",
    )
    client = otp_client(regulator_admin)

    with patch("securityobjectives.views.celery_health_check", return_value=False):
        response = client.get(reverse("security_objectives_export_status", args=[export.id]))

    assert response.json()["status"] == "FAIL"
    export.refresh_from_db()
    assert export.task_status == "FAIL"


@pytest.mark.django_db
def test_status_leaves_a_running_job_alone_when_the_stack_is_healthy(otp_client, populate_so_db):
    regulator_admin = get_user(populate_so_db, "regadmin@reg1.lu")
    grant_export_right(regulator_admin)
    export = SecurityObjectiveExport.objects.create(
        user=regulator_admin,
        regulation_id=populate_so_db["so_standard"][0].regulation_id,
        filename="SO_export.xlsx",
    )
    client = otp_client(regulator_admin)

    with patch("securityobjectives.views.celery_health_check", return_value=True):
        response = client.get(reverse("security_objectives_export_status", args=[export.id]))

    assert response.json()["status"] == "RUNNING"


@pytest.mark.django_db
def test_export_view_rejects_an_empty_multiselect(otp_client, populate_so_db):
    regulator_admin = get_user(populate_so_db, "regadmin@reg1.lu")
    grant_export_right(regulator_admin)
    client = otp_client(regulator_admin)
    filters = build_filters(populate_so_db)
    filters["sectors"] = []

    with patch("securityobjectives.views.generate_so_export_task.delay") as delay:
        response = client.post(reverse("export_security_objectives"), filters)

    assert response.status_code == 400
    assert not SecurityObjectiveExport.objects.exists()
    delay.assert_not_called()


@pytest.mark.django_db
def test_xlsx_export_holds_the_dashboard_values(populate_so_db, settings, tmp_path):
    regulator_admin = get_user(populate_so_db, "regadmin@reg1.lu")
    grant_export_right(regulator_admin)
    declaration = populate_so_db["sas"][0]

    export = run_export(regulator_admin, build_filters(populate_so_db), settings, tmp_path)

    assert export.task_status == "DONE"
    workbook = load_export_workbook(export)
    assert len(workbook.sheetnames) == 2

    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.values)
    header, row = rows[0], rows[1]
    assert header[0] == "Status"
    assert row[0] == declaration.get_status_display()
    assert row[3] == declaration.group.group_id
    assert row[4] == declaration.standard.label
    assert row[5] == declaration.creator_company_name
    assert row[7] == declaration.year_of_submission
    assert row[8].endswith("%")


@pytest.mark.django_db
def test_metadata_sheet_records_the_row_count(populate_so_db, settings, tmp_path):
    regulator_admin = get_user(populate_so_db, "regadmin@reg1.lu")
    grant_export_right(regulator_admin)

    export = run_export(regulator_admin, build_filters(populate_so_db), settings, tmp_path)

    workbook = load_export_workbook(export)
    metadata = dict(workbook[workbook.sheetnames[1]].values)
    assert metadata["Number of rows"] == "1"
    assert metadata["Author"] == f"{regulator_admin.get_full_name()} ({regulator_admin.email})"


@pytest.mark.django_db
def test_csv_export_is_a_zip_of_two_files(populate_so_db, settings, tmp_path):
    regulator_admin = get_user(populate_so_db, "regadmin@reg1.lu")
    grant_export_right(regulator_admin)

    export = run_export(regulator_admin, build_filters(populate_so_db, file_format="csv"), settings, tmp_path)

    with zipfile.ZipFile(export.get_file_path()) as archive:
        assert sorted(archive.namelist()) == ["declarations.csv", "metadata.csv"]


@pytest.mark.django_db
def test_export_logs_every_declaration_and_sends_one_email(
    populate_so_db,
    settings,
    tmp_path,
    create_standard_answer_group,
    create_standard_answer,
):
    regulator_admin = get_user(populate_so_db, "regadmin@reg1.lu")
    grant_export_right(regulator_admin)
    first = populate_so_db["sas"][0]
    second = create_standard_answer(
        standard=first.standard,
        submitter_user=first.submitter_user,
        submitter_company=first.submitter_company,
        creator_name=first.creator_name,
        creator_company_name=first.creator_company_name,
        sectors=list(first.sectors.all()),
        group=create_standard_answer_group(company=first.submitter_company),
        status=first.status,
        year_of_submission=first.year_of_submission,
    )

    run_export(regulator_admin, build_filters(populate_so_db), settings, tmp_path)

    entries = LogEntry.objects.filter(action_flag=EXPORT)
    assert {entry.object_id for entry in entries} == {str(first.id), str(second.id)}
    for entry in entries:
        assert "A total of 2 security objective declarations were exported" in entry.change_message
    # The mail is per export, not per declaration.
    assert len(mail.outbox) == 1


@pytest.mark.django_db
def test_a_failed_export_writes_no_log_and_no_email(populate_so_db, settings, tmp_path):
    regulator_admin = get_user(populate_so_db, "regadmin@reg1.lu")
    grant_export_right(regulator_admin)
    settings.PATH_FOR_REPORTING_PDF = str(tmp_path)
    export = SecurityObjectiveExport.objects.create(
        user=regulator_admin,
        regulation_id=populate_so_db["so_standard"][0].regulation_id,
        filename="SO_export.xlsx",
    )

    with patch("securityobjectives.tasks.write_export_xlsx", side_effect=OSError("disk full")):
        with pytest.raises(OSError):
            generate_so_export_task(export.id, regulator_admin.id, build_filters(populate_so_db), "en")

    export.refresh_from_db()
    assert export.task_status == "FAIL"
    assert not LogEntry.objects.filter(action_flag=EXPORT).exists()
    assert len(mail.outbox) == 0


@pytest.mark.django_db
def test_a_sector_outside_the_scope_returns_no_rows(populate_so_db, settings, tmp_path):
    regulator_user = get_user(populate_so_db, "reguser@reg1.lu")
    grant_export_right(regulator_user)
    declaration = populate_so_db["sas"][0]
    out_of_scope = selectable_sectors(declaration).first()
    in_scope = Sector.objects.filter(parent__isnull=False).exclude(id=out_of_scope.id).first()
    RegulatorUser.objects.get(user=regulator_user, regulator=regulator_user.regulators.first()).sectors.set([in_scope])

    filters = build_filters(populate_so_db, sectors=[str(out_of_scope.id)])
    export = run_export(regulator_user, filters, settings, tmp_path)

    workbook = load_export_workbook(export)
    assert len(list(workbook[workbook.sheetnames[0]].values)) == 1


@pytest.mark.django_db
def test_cleanup_removes_the_file_and_the_job(populate_so_db, settings, tmp_path):
    regulator_admin = get_user(populate_so_db, "regadmin@reg1.lu")
    grant_export_right(regulator_admin)
    export = run_export(regulator_admin, build_filters(populate_so_db), settings, tmp_path)
    file_path = Path(export.get_file_path())
    assert file_path.exists()

    cleanup_so_export_file(export.id)

    assert not file_path.exists()
    assert not SecurityObjectiveExport.objects.filter(id=export.id).exists()


@pytest.mark.django_db
def test_cleanup_tolerates_an_already_deleted_file(populate_so_db, settings, tmp_path):
    regulator_admin = get_user(populate_so_db, "regadmin@reg1.lu")
    grant_export_right(regulator_admin)
    export = run_export(regulator_admin, build_filters(populate_so_db), settings, tmp_path)
    Path(export.get_file_path()).unlink()

    cleanup_so_export_file(export.id)

    assert not SecurityObjectiveExport.objects.filter(id=export.id).exists()


@pytest.mark.django_db
def test_cleanup_of_an_unknown_job_is_a_no_op():
    cleanup_so_export_file(404)


@pytest.mark.django_db
def test_download_is_refused_to_another_user(otp_client, populate_so_db, settings, tmp_path):
    owner = get_user(populate_so_db, "regadmin@reg1.lu")
    grant_export_right(owner)
    other = get_user(populate_so_db, "regadmin@reg2.lu")
    grant_export_right(other)
    export = run_export(owner, build_filters(populate_so_db), settings, tmp_path)

    client = otp_client(other)
    response = client.get(reverse("download_security_objectives_export", args=[export.file_uuid]))

    assert response.status_code == 404


@pytest.mark.django_db
def test_owner_can_download_the_export(otp_client, populate_so_db, settings, tmp_path):
    owner = get_user(populate_so_db, "regadmin@reg1.lu")
    grant_export_right(owner)
    export = run_export(owner, build_filters(populate_so_db), settings, tmp_path)

    client = otp_client(owner)
    response = client.get(reverse("download_security_objectives_export", args=[export.file_uuid]))

    assert response.status_code == 200
    assert response["Content-Disposition"].startswith("attachment;")
